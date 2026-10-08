"""
Cover Page and Table of Contents (TOC) Generation Engine.
Implements surgical template modification and standalone academic cover page creation.
"""

from typing import Optional, List
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .config import DocumentMetadata, TypographyConfig
from .bidi import set_paragraph_bidi, style_run, tokenize_runs, apply_paragraph_spacing


def update_cover_text_node(p_element, new_text: str):
    """
    Surgically modifies <w:t> text nodes directly inside an element
    to preserve underlying Complex Script, font, and layout properties.
    """
    t_nodes = p_element.xpath('.//w:t')
    if t_nodes:
        t_nodes[0].text = new_text
        for extra_t in t_nodes[1:]:
            extra_t.getparent().remove(extra_t)


def prepare_template_document(doc, meta: DocumentMetadata):
    """
    Surgically cleans old body elements from template while preserving:
    - Page 1: Vector borders & Cover page skeleton
    - Page 2: TOC Heading & <w:sdt> field code
    - Section properties & page geometry (2.0 cm margins)
    """
    body = doc._body._element
    children = list(body)
    sectPr = body.find(qn('w:sectPr'))

    # Surgical text replacement on Cover Page
    if meta.title and len(children) > 1:
        update_cover_text_node(children[1], meta.title)
    if meta.subtitle and len(children) > 2:
        update_cover_text_node(children[2], meta.subtitle)
    if meta.supervisors and len(children) > 4:
        sup_str = ("استاد راهنما: " if len(meta.supervisors) == 1 else "اساتید: ") + "، ".join(meta.supervisors)
        update_cover_text_node(children[4], sup_str)
    if meta.authors and len(children) > 5:
        if any("دانشجو" in a for a in meta.authors):
            auth_str = "دانشجو: " + "، ".join(meta.authors)
        elif len(meta.authors) == 1:
            auth_str = "نگارنده: " + "، ".join(meta.authors)
        else:
            auth_str = "نگارندگان: " + "، ".join(meta.authors)
        update_cover_text_node(children[5], auth_str)
    if meta.date and len(children) > 7:
        update_cover_text_node(children[7], meta.date)
    if meta.toc_title and len(children) > 9:
        update_cover_text_node(children[9], meta.toc_title)

    # In standard template:
    # 0..7: Cover page elements
    # 8: Page break (Page 1 -> Page 2)
    # 9: TOC heading ("فهرست مطالب")
    # 10: <w:sdt> TOC field
    # 11: Empty spacer
    # 12: Page break (Page 2 -> Page 3)
    # 13+: Old body content
    body_start_index = 13 if meta.toc else 9
    if not meta.toc and len(children) > 9:
        # If user explicitly opted out of TOC, remove TOC elements 9..12
        for child in children[9:13]:
            if child in body:
                body.remove(child)
        body_start_index = 9

    for child in children[body_start_index:]:
        if child != sectPr and child in body:
            body.remove(child)


def build_standalone_cover_page(doc, meta: DocumentMetadata, config: TypographyConfig):
    """
    Builds a pristine standalone academic cover page from scratch.
    """
    # Spacer
    sp = doc.add_paragraph()
    apply_paragraph_spacing(sp, before_pt=40.0, after_pt=20.0)

    # Institution
    if meta.institution:
        p_inst = doc.add_paragraph()
        p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_bidi(p_inst, is_rtl=(meta.lang == "fa"), justify=False)
        apply_paragraph_spacing(p_inst, before_pt=0, after_pt=8, line_spacing_multiplier=1.3)
        r = p_inst.add_run(meta.institution)
        style_run(r, font_fa=config.font_fa_heading, font_en=config.font_en_heading,
                  size_pt=15.0, is_fa=(meta.lang == "fa"), bold=True)

    if meta.faculty:
        p_fac = doc.add_paragraph()
        p_fac.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_bidi(p_fac, is_rtl=(meta.lang == "fa"), justify=False)
        apply_paragraph_spacing(p_fac, before_pt=0, after_pt=30, line_spacing_multiplier=1.3)
        r = p_fac.add_run(meta.faculty)
        style_run(r, font_fa=config.font_fa_body, font_en=config.font_en_body,
                  size_pt=13.0, is_fa=(meta.lang == "fa"), bold=False)

    # Title
    if meta.title:
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_bidi(p_title, is_rtl=(meta.lang == "fa"), justify=False)
        apply_paragraph_spacing(p_title, before_pt=30, after_pt=12, line_spacing_multiplier=1.4)
        for seg_type, seg_text in tokenize_runs(meta.title):
            r = p_title.add_run(seg_text)
            style_run(r, font_fa=config.font_fa_heading, font_en=config.font_en_heading,
                      size_pt=config.size_title, is_fa=(seg_type == 'fa'), bold=True)

    # Subtitle
    if meta.subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_bidi(p_sub, is_rtl=(meta.lang == "fa"), justify=False)
        apply_paragraph_spacing(p_sub, before_pt=0, after_pt=40, line_spacing_multiplier=1.4)
        for seg_type, seg_text in tokenize_runs(meta.subtitle):
            r = p_sub.add_run(seg_text)
            style_run(r, font_fa=config.font_fa_heading, font_en=config.font_en_heading,
                      size_pt=config.size_subtitle, is_fa=(seg_type == 'fa'), bold=True)

    # Metadata Block (Supervisors, Authors)
    p_spacer = doc.add_paragraph()
    apply_paragraph_spacing(p_spacer, before_pt=30, after_pt=20)

    if meta.supervisors:
        p_sup = doc.add_paragraph()
        p_sup.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_bidi(p_sup, is_rtl=(meta.lang == "fa"), justify=False)
        apply_paragraph_spacing(p_sup, before_pt=0, after_pt=8)
        lbl = "اساتید راهنما: " if meta.lang == "fa" else "Supervisors: "
        val = "، ".join(meta.supervisors) if meta.lang == "fa" else ", ".join(meta.supervisors)
        r = p_sup.add_run(f"{lbl}{val}")
        style_run(r, font_fa=config.font_fa_body, font_en=config.font_en_body,
                  size_pt=14.0, is_fa=(meta.lang == "fa"), bold=False)

    if meta.authors:
        p_auth = doc.add_paragraph()
        p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_bidi(p_auth, is_rtl=(meta.lang == "fa"), justify=False)
        apply_paragraph_spacing(p_auth, before_pt=0, after_pt=8)
        lbl = "نگارش: " if meta.lang == "fa" else "Authors: "
        val = "، ".join(meta.authors) if meta.lang == "fa" else ", ".join(meta.authors)
        r = p_auth.add_run(f"{lbl}{val}")
        style_run(r, font_fa=config.font_fa_body, font_en=config.font_en_body,
                  size_pt=14.0, is_fa=(meta.lang == "fa"), bold=True)

    if meta.date:
        p_date = doc.add_paragraph()
        p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_bidi(p_date, is_rtl=(meta.lang == "fa"), justify=False)
        apply_paragraph_spacing(p_date, before_pt=30, after_pt=0)
        r = p_date.add_run(meta.date)
        style_run(r, font_fa=config.font_fa_body, font_en=config.font_en_body,
                  size_pt=14.0, is_fa=(meta.lang == "fa"), bold=True)

    # Page Break after Cover Page
    doc.add_page_break()


def insert_standalone_toc(doc, toc_title: str = "فهرست مطالب", config: Optional[TypographyConfig] = None):
    """
    Inserts a native Word <w:sdt> Table of Contents element with page break.
    """
    cfg = config or TypographyConfig()
    
    # TOC Heading
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_bidi(p_head, is_rtl=True, justify=False)
    apply_paragraph_spacing(p_head, before_pt=14.0, after_pt=14.0)
    r = p_head.add_run(toc_title)
    style_run(r, font_fa=cfg.font_fa_heading, font_en=cfg.font_en_heading,
              size_pt=16.0, is_fa=True, bold=True)

    # Native Word TOC SDT XML
    toc_xml = r'''
    <w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:sdtPr>
        <w:docPartObj>
          <w:docPartGallery w:val="Table of Contents"/>
          <w:docPartUnique/>
        </w:docPartObj>
      </w:sdtPr>
      <w:sdtContent>
        <w:p>
          <w:pPr>
            <w:pStyle w:val="TOC1"/>
            <w:bidi/>
            <w:jc w:val="both"/>
          </w:pPr>
          <w:r>
            <w:fldChar w:fldCharType="begin"/>
            <w:instrText xml:space="preserve"> TOC \o "1-3" \h \z \u </w:instrText>
            <w:fldChar w:fldCharType="separate"/>
            <w:fldChar w:fldCharType="end"/>
          </w:r>
        </w:p>
      </w:sdtContent>
    </w:sdt>
    '''
    sdt_elem = parse_xml(toc_xml)
    doc._body._element.append(sdt_elem)

    # Page break after TOC
    doc.add_page_break()
