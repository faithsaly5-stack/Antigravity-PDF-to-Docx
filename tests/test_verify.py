"""
Comprehensive verification test script for Persian academic DOCX generation.
"""

import os
import sys
import docx
from docx.oxml.ns import qn

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from md2docx import MarkdownToDocx

def test_document():
    doc_path = os.path.join(PROJECT_ROOT, "output_academic.docx")
    sample_path = os.path.join(PROJECT_ROOT, "samples", "sample_persian_academic.md")
    template_path = os.path.join(PROJECT_ROOT, "templates", "default_academic.docx")

    if not os.path.exists(doc_path):
        print(f"Generating {doc_path} from {sample_path}...")
        conv = MarkdownToDocx(default_template=template_path)
        conv.convert_file(sample_path, doc_path, template=template_path, update_toc=False, export_pdf=False)

    doc = docx.Document(doc_path)
    print(f"Document top-level paragraphs: {len(doc.paragraphs)}")
    print(f"Document tables: {len(doc.tables)}")

    # 1. Verify Cover Page
    p1 = doc.paragraphs[1]
    print(f"Cover Title: '{p1.text}'")
    assert "گزارش کار جامع آزمایشگاه شیمی عمومی و تجزیه" in p1.text

    # Collect all paragraphs including table cells
    all_paras = list(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                all_paras.extend(cell.paragraphs)

    # 2. Verify Mixed-Script Run Tokenization
    found_mixed = False
    for p in all_paras:
        # Check for paragraph discussing ion balance or callout
        if "اثر یون مشترک" in p.text or "خورنده هستند" in p.text or "قانون بیر-لامبرت" in p.text:
            found_mixed = True
            print(f"\n[Mixed Script Paragraph]: {p.text}")
            for r in p.runs:
                rPr = r._element.find(qn('w:rPr'))
                rtl = rPr.find(qn('w:rtl')) is not None if rPr is not None else False
                rFonts = rPr.find(qn('w:rFonts')) if rPr is not None else None
                ascii_f = rFonts.get(qn('w:ascii')) if rFonts is not None else None
                cs_f = rFonts.get(qn('w:cs')) if rFonts is not None else None
                print(f"   Run: rtl={str(rtl):5} | font_ascii={ascii_f} | font_cs={cs_f} | text='{r.text}'")
                
                # Check Latin tokens
                if any(x in r.text for x in ["Fe", "SCN", "mol", "cm", "nm", "HCl", "NaOH"]):
                    assert not rtl, f"Latin run '{r.text}' must NOT have w:rtl!"

    assert found_mixed, "Mixed script paragraph not found!"

    # 3. Verify Reference BIDI Isolation
    refs = [p for p in doc.paragraphs if any(k in p.text for k in ["Chang, R.", "Skoog, D.", "زارع، محمدرضا", "افشار، احمد"])]
    print(f"\n[References]: found {len(refs)}")
    for p in refs:
        pPr = p._element.find(qn('w:pPr'))
        bidi = pPr.find(qn('w:bidi')) is not None if pPr is not None else False
        print(f"   Ref bidi={str(bidi):5} | text='{p.text[:60]}...'")
        if "Chang" in p.text or "Skoog" in p.text:
            assert not bidi, f"English ref '{p.text[:30]}' must NOT have w:bidi!"
        else:
            assert bidi, f"Persian ref '{p.text[:30]}' MUST have w:bidi!"

    # 4. Verify Math Elements (both top-level and in tables)
    math_count = 0
    for p in all_paras:
        omaths = p._element.xpath('.//m:oMath | .//m:oMathPara')
        if omaths:
            math_count += len(omaths)
    print(f"\n[Math Equations]: Found {math_count} native OMML equations across all paragraphs and cells")
    assert math_count >= 2, f"Expected at least 2 OMML equations, found {math_count}"

    # 5. Verify Table
    data_tbl = None
    for tbl in doc.tables:
        if len(tbl.rows) >= 5 and len(tbl.columns) >= 4:
            data_tbl = tbl
            break
    assert data_tbl is not None, "Academic data table not found!"
    bidi_vis = data_tbl._element.xpath('.//w:bidiVisual')
    print(f"\n[Data Table]: rows={len(data_tbl.rows)}, cols={len(data_tbl.columns)}, bidiVisual={bool(bidi_vis)}")
    assert bool(bidi_vis), "Persian academic table MUST have w:bidiVisual!"

    print("\n[ALL AUDIT CHECKS PASSED PERFECTLY!]")

if __name__ == "__main__":
    test_document()
